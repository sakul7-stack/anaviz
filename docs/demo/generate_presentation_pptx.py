#!/usr/bin/env python3
"""Generate a presentation-ready Anaviz PowerPoint.

Documentation artifact only: this script does not import or run Anaviz.
"""
from pathlib import Path

from PIL import Image as PILImage
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
IMG = ROOT / "images"
OUT = ROOT / "ANAVIZ_PRESENTATION.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Theme colors
NAVY = RGBColor(20, 54, 92)
BLUE = RGBColor(31, 111, 139)
BLUE_LIGHT = RGBColor(225, 239, 252)
GREEN = RGBColor(25, 135, 84)
GREEN_LIGHT = RGBColor(225, 244, 235)
PURPLE = RGBColor(111, 66, 193)
PURPLE_LIGHT = RGBColor(239, 232, 250)
ORANGE = RGBColor(217, 107, 0)
ORANGE_LIGHT = RGBColor(255, 239, 220)
RED = RGBColor(198, 40, 40)
RED_LIGHT = RGBColor(252, 231, 232)
INK = RGBColor(25, 32, 42)
MUTED = RGBColor(88, 101, 116)
WHITE = RGBColor(255, 255, 255)
CODE_BG = RGBColor(27, 38, 54)
CODE_TEXT = RGBColor(229, 239, 250)
GOLD = RGBColor(138, 75, 8)

W = 13.333
H = 7.5


def add_shape(slide, left, top, width, height, fill, line=None, radius=True):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    shape.line.width = Pt(1)
    return shape


def add_text(slide, text, left, top, width, height, size=18, color=INK,
             bold=False, font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rich_lines(slide, lines, left, top, width, height, size=18, color=INK,
                   bullet=False, line_spacing=1.05):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    for index, item in enumerate(lines):
        text = item[0] if isinstance(item, tuple) else item
        item_color = item[1] if isinstance(item, tuple) and len(item) > 1 else color
        item_bold = item[2] if isinstance(item, tuple) and len(item) > 2 else False
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(7)
        p.line_spacing = line_spacing
        if bullet:
            p.text = "• " + text
        else:
            p.text = text
        for run in p.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(size)
            run.font.color.rgb = item_color
            run.font.bold = item_bold
    return box


def add_title(slide, title, subtitle=None, section=None):
    add_shape(slide, 0, 0, W, 0.16, BLUE, BLUE, radius=False)
    add_text(slide, title, 0.55, 0.35, 12.1, 0.55, size=27, color=NAVY, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.58, 0.93, 12.0, 0.35, size=12, color=MUTED)
    if section:
        add_text(slide, section.upper(), 11.3, 0.42, 1.4, 0.3, size=9, color=BLUE, bold=True, align=PP_ALIGN.RIGHT)


def add_footer(slide, number):
    add_shape(slide, 0.55, 7.16, 12.25, 0.012, RGBColor(222, 228, 235), RGBColor(222, 228, 235), radius=False)
    add_text(slide, "ANAVIZ  |  Dataset-independent time-series visualization", 0.58, 7.22, 8.5, 0.18, size=8, color=MUTED)
    add_text(slide, str(number), 12.2, 7.22, 0.55, 0.18, size=8, color=MUTED, align=PP_ALIGN.RIGHT)


def add_cue(slide, text, top=6.55):
    add_shape(slide, 0.55, top, 12.25, 0.42, ORANGE_LIGHT, ORANGE, radius=True)
    add_text(slide, "SPEAKER CUE  |  " + text, 0.72, top + 0.07, 11.9, 0.28, size=11, color=GOLD, bold=True, valign=MSO_ANCHOR.MIDDLE)


def add_code(slide, source, code, left, top, width, height, explanation=None):
    add_shape(slide, left, top, width, height, CODE_BG, CODE_BG, radius=True)
    add_text(slide, source, left + 0.18, top + 0.10, width - 0.35, 0.25, size=9, color=RGBColor(150, 202, 255), bold=True, font="Aptos")
    box = slide.shapes.add_textbox(Inches(left + 0.18), Inches(top + 0.42), Inches(width - 0.35), Inches(height - 0.52))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    for index, line in enumerate(code.strip("\n").split("\n")):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = line
        p.space_after = Pt(0)
        p.line_spacing = 0.88
        for run in p.runs:
            run.font.name = "Cascadia Mono"
            run.font.size = Pt(10)
            run.font.color.rgb = CODE_TEXT
    if explanation:
        add_text(slide, "Plain English: " + explanation, left, top + height + 0.08, width, 0.55, size=12, color=GREEN, bold=True)


def add_image_fit(slide, filename, left, top, width, height):
    path = IMG / filename
    if not path.exists():
        add_text(slide, "Missing image: " + filename, left, top, width, height, size=14, color=RED)
        return
    with PILImage.open(path) as im:
        iw, ih = im.size
    image_ratio = iw / ih
    box_ratio = width / height
    if image_ratio > box_ratio:
        draw_width = width
        draw_height = width / image_ratio
        x = left
        y = top + (height - draw_height) / 2
    else:
        draw_height = height
        draw_width = height * image_ratio
        x = left + (width - draw_width) / 2
        y = top
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(draw_width), height=Inches(draw_height))


def new_slide(title, subtitle=None, section=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, title, subtitle, section)
    return slide


# 1 — Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_shape(slide, 0, 0, W, H, NAVY, NAVY, radius=False)
add_shape(slide, 0.62, 0.72, 0.16, 5.95, ORANGE, ORANGE, radius=False)
add_text(slide, "ANAVIZ", 1.12, 1.15, 10.8, 0.9, size=46, color=WHITE, bold=True)
add_text(slide, "Dataset-independent time-series visualization", 1.15, 2.12, 10.7, 0.6, size=25, color=RGBColor(205, 225, 247), bold=True)
add_text(slide, "How the code connects APIs, caches data, preserves gaps, and renders evidence", 1.17, 2.9, 10.5, 0.55, size=17, color=RGBColor(222, 232, 242))
add_shape(slide, 1.15, 4.2, 5.15, 1.15, RGBColor(34, 77, 119), RGBColor(72, 125, 178), radius=True)
add_text(slide, "Example source: CERN ATLAS HLT p-beat\nProject: generic HTTP time-series explorer", 1.38, 4.48, 4.7, 0.55, size=16, color=WHITE, bold=True)
add_text(slide, "Presentation-ready deck  |  code + diagrams + speaking cues", 1.15, 6.7, 10.8, 0.28, size=11, color=RGBColor(190, 211, 233))

# 2 — Big idea
slide = new_slide("The big idea", "Anaviz is an adapter and evidence layer, not an HLT-only chart.", "Overview")
add_shape(slide, 0.65, 1.55, 3.55, 3.7, RED_LIGHT, RED, radius=True)
add_text(slide, "ANY HTTP\nTIME-SERIES API", 0.95, 2.05, 2.95, 0.75, size=24, color=RED, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "timestamps + values\npossibly different field names\npossibly different JSON shape", 0.95, 3.1, 2.95, 0.9, size=15, color=INK, align=PP_ALIGN.CENTER)
add_shape(slide, 4.9, 1.55, 3.55, 3.7, ORANGE_LIGHT, ORANGE, radius=True)
add_text(slide, "config.json\n+ adapter", 5.25, 2.05, 2.85, 0.75, size=24, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "maps source fields\nfetches uncovered ranges\ntranslates to canonical models", 5.25, 3.1, 2.85, 0.9, size=15, color=INK, align=PP_ALIGN.CENTER)
add_shape(slide, 9.15, 1.55, 3.55, 3.7, GREEN_LIGHT, GREEN, radius=True)
add_text(slide, "CHART\n+ HEATMAP\n+ HISTOGRAM", 9.45, 2.0, 2.95, 1.15, size=24, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "fast display\nvisible gaps\nquality + fidelity", 9.45, 3.45, 2.95, 0.75, size=15, color=INK, align=PP_ALIGN.CENTER)
add_text(slide, "→", 4.25, 2.8, 0.55, 0.6, size=38, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "→", 8.45, 2.8, 0.55, 0.6, size=38, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
add_cue(slide, "The HLT dataset demonstrates the architecture; the frontend is designed for any timestamp-and-value HTTP API.")
add_footer(slide, 2)

# 3 — Dataset
slide = new_slide("The example dataset: ATLAS HLT p-beat", "Real monitoring data from CERN's ATLAS trigger system.", "Dataset")
add_rich_lines(slide, [
    ("HLT = High Level Trigger", NAVY, True),
    ("The system makes fast event-selection decisions at the LHC.", INK, False),
    ("p-beat = periodic heartbeat/rate monitoring", NAVY, True),
    ("Each channel represents a DCM computing node and its measured rate.", INK, False),
], 0.75, 1.35, 5.55, 2.1, size=18)
add_shape(slide, 0.75, 3.7, 5.55, 1.55, BLUE_LIGHT, BLUE, radius=True)
add_rich_lines(slide, [
    ("Data facts", BLUE, True),
    ("~3,690 channels  |  large CSV files  |  long time range", INK, False),
    ("NaN means no valid sample — not zero.", RED, True),
], 0.95, 3.92, 5.15, 1.1, size=16)
add_code(slide, "datasource/ingest_hlt.py", "ELEMENT_ID_OFFSET = 900_000\n\n# column 0 -> 900000\n# column 1 -> 900001", 6.7, 1.45, 5.75, 2.0, "The importer gives each source column a stable entity ID.")
add_code(slide, "datasource/ingest_hlt.py", "chunksize = 5000\nreader = pd.read_csv(path, chunksize=chunksize)\n\n# process a small part of the large CSV", 6.7, 4.0, 5.75, 1.55, "Chunking prevents a multi-gigabyte file from filling memory.")
add_cue(slide, "Introduce the data, then immediately connect it to the design problem: large volume plus missing samples.")
add_footer(slide, 3)

# 4 — Architecture image
slide = new_slide("Architecture: two independent containers", "Separate ownership protects the real dataset.", "Architecture")
add_image_fit(slide, "06_architecture.png", 0.65, 1.25, 12.05, 4.95)
add_cue(slide, "Datasource owns the source of truth. Project owns generic visualization and cache. They communicate over HTTP, not shared database tables.")
add_footer(slide, 4)

# 5 — API contract image
slide = new_slide("The bridge: source API → config → canonical contract", "The frontend does not learn every source's database vocabulary.", "Architecture")
add_image_fit(slide, "09_api_contract_translation.png", 0.7, 1.25, 11.95, 4.9)
add_cue(slide, "The adapter translates source-specific names like element_id and quality_flag into generic entities, points, gaps, fidelity, and metrics.")
add_footer(slide, 5)

# 6 — Config code
slide = new_slide("config.json: declarative integration", "The user describes the API instead of writing a new renderer.", "Configuration")
add_code(slide, "user config.json", "\"source\": {\"base_url\": \"http://datasource:9000\"}\n\"series\": {\n  \"path\": \"/archive/eventhistory\",\n  \"shape\": \"columnar\",\n  \"parameters\": {\n    \"entity\": \"element_id\", \"start\": \"t0\", \"end\": \"t1\"\n  }\n}", 0.7, 1.35, 6.1, 3.9, "This maps Anaviz names to the source API's paths and parameters.")
add_rich_lines(slide, [
    ("Important fields", BLUE, True),
    ("source.base_url — where to fetch", INK, False),
    ("endpoints — which paths to call", INK, False),
    ("mapping — where timestamp/value/quality live", INK, False),
    ("pagination — how to follow pages", INK, False),
    ("rollup_levels — optional bucket sizes", INK, False),
], 7.25, 1.45, 5.0, 2.8, size=17)
add_code(slide, "server/api/config.py", "if not value.startswith((\"http://\", \"https://\")):\n    raise ValueError(\"base_url must be http(s)\")", 7.25, 4.4, 5.0, 1.2, "Validation rejects unsafe or unusable source URLs before registration.")
add_cue(slide, "config.json is the bridge. It is data, not executable code.")
add_footer(slide, 6)

# 7 — Canonical contract
slide = new_slide("Canonical request and response", "The contract keeps backend and frontend independent.", "API")
add_shape(slide, 0.7, 1.35, 5.85, 4.65, BLUE_LIGHT, BLUE, radius=True)
add_text(slide, "REQUEST: SeriesQuery", 0.95, 1.65, 5.3, 0.35, size=20, color=BLUE, bold=True)
add_code(slide, "server/api/contracts.py", "dataset_id\nentity_ids\nmeasure_ids\nrange: start / end\nresolution: pixels × points\ndownsampling: M4 / LTTB / RAW", 0.95, 2.15, 5.25, 2.4, "The browser describes what it wants; it does not choose SQL.")
add_shape(slide, 6.85, 1.35, 5.85, 4.65, GREEN_LIGHT, GREEN, radius=True)
add_text(slide, "RESPONSE: CanonicalSeries", 7.1, 1.65, 5.3, 0.35, size=20, color=GREEN, bold=True)
add_code(slide, "server/api/contracts.py", "points: t / value / min / max\ngaps\nquality_summary\nfidelity\nmetrics: source / scanned / returned", 7.1, 2.15, 5.25, 2.4, "The result explains both the values and how they were produced.")
add_cue(slide, "This is the central abstraction: source-specific data enters, canonical evidence leaves.")
add_footer(slide, 7)

# 8 — Query pipeline
slide = new_slide("Query pipeline: from click to chart", "One user action triggers a controlled six-stage flow.", "Query")
add_image_fit(slide, "07_query_pipeline.png", 0.7, 1.35, 12.0, 4.55)
add_cue(slide, "Hydrate only missing ranges, select a suitable tier, downsample to the pixel budget, and return fidelity metadata.")
add_footer(slide, 8)

# 9 — Hydration/cache
slide = new_slide("Hydration: fetch only what is missing", "The first query may be slower; later queries reuse the project cache.", "Cache")
add_code(slide, "server/api/configurable.py", "covered = await cache.coverage(entity, measure, start, end)\ngaps = coverage_gaps(start, end, covered)\n\nfor a, b in gaps:\n    rows = await fetch_source(a, b)\n    await cache.store(entity, measure, rows)", 0.7, 1.35, 6.3, 3.75, "Coverage tells the adapter which portions have already been checked. Only uncovered portions go back to the source.")
add_image_fit(slide, "10_cache_layers.png", 7.25, 1.3, 5.55, 4.45)
add_cue(slide, "cache_series stores samples; cache_coverage avoids repeated probes; cache_rollup stores optional summaries. dataset_key isolates configurations.")
add_footer(slide, 9)

# 10 — Rollups and tiers
slide = new_slide("Rollup tiers: less data, same important evidence", "Rollups are explicit and opt-in.", "Performance")
add_image_fit(slide, "03_rollup_buckets.png", 0.65, 1.28, 6.0, 4.8)
add_image_fit(slide, "04_tier_selection.png", 6.95, 1.28, 5.8, 4.8)
add_cue(slide, "A bucket keeps first, min, max, last, average, count, and quality. Short spans stay raw; large spans use coarser configured tiers.")
add_footer(slide, 10)

# 11 — Downsampling
slide = new_slide("Downsampling: fit the screen, preserve the shape", "The point budget comes from pixels, not an arbitrary number.", "Performance")
add_image_fit(slide, "02_downsampling.png", 0.65, 1.12, 7.05, 5.25)
add_shape(slide, 8.05, 1.35, 4.55, 3.6, PURPLE_LIGHT, PURPLE, radius=True)
add_rich_lines(slide, [
    ("Algorithms", PURPLE, True),
    ("RAW — exact selected points", INK, False),
    ("M4 — first/min/max/last", INK, False),
    ("LTTB — shape-aware selection", INK, False),
    ("MINMAXLTTB — shape + extrema", INK, False),
], 8.35, 1.7, 3.95, 1.95, size=16)
add_code(slide, "server/downsample/__init__.py", "budget = pixel_width * points_per_pixel\n\nif algo == \"RAW\": return payload\nif algo == \"M4\": return m4_series(...)\nif algo == \"LTTB\": return lttb_indices(...)", 8.05, 4.2, 4.55, 1.3, "Only the display response is reduced; raw cache rows remain available.")
add_cue(slide, "With 1,000 pixels and 2 points per pixel, the target is about 2,000 display points.")
add_footer(slide, 11)

# 12 — Gaps
slide = new_slide("Gap preservation: missing data is evidence", "Anaviz does not draw a false line through an unobserved interval.", "Fidelity")
add_image_fit(slide, "01_raw_series.png", 0.65, 1.2, 5.9, 2.45)
add_image_fit(slide, "05_gap_preservation.png", 6.75, 1.2, 5.95, 4.85)
add_code(slide, "server/api/common.py + frontend/utils.js", "if timestamp_gap > threshold:\n    gaps.append([start, end])\n\nif in_gap:\n    value = null", 0.75, 4.05, 5.65, 1.2, "The backend declares the gap; the frontend inserts null so the line breaks.")
add_cue(slide, "Missing is not zero. The tooltip and visible break tell the operator that the source had no observation.")
add_footer(slide, 12)

# 13 — No data loss
slide = new_slide("Dense ranges: no silent truncation", "row_cap triggers aggregation, not an early chart stop.", "Fidelity")
add_image_fit(slide, "08_no_data_loss.png", 0.7, 1.18, 7.15, 4.95)
add_code(slide, "server/api/configurable.py — dense path", "if total > row_cap:\n    bucket_s = span_seconds / pixel_budget\n    data = read_bucketed(bucket_s)\n\n# SQL uses min, max, avg, count, first, last", 8.15, 1.6, 4.45, 2.5, "Every cached row contributes to a pixel-sized bucket; the rendered response still preserves the full time range and extrema.")
add_shape(slide, 8.15, 4.55, 4.45, 1.0, GREEN_LIGHT, GREEN, radius=True)
add_text(slide, "Raw rows remain in cache_series.\nThe chart receives a faithful summary, not a truncated prefix.", 8.4, 4.78, 3.95, 0.52, size=15, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
add_cue(slide, "The important distinction: not every raw row is drawn individually, but no part of the requested range is silently thrown away.")
add_footer(slide, 13)

# 14 — Frontend
slide = new_slide("One response, three coordinated frontend views", "The canonical response feeds chart, heatmap, and histogram.", "Frontend")
add_image_fit(slide, "11_frontend_panels.png", 0.65, 1.25, 12.05, 4.85)
add_cue(slide, "chart.js shows time and bands; heatmap.js shows abnormality by entity and time; histogram.js shows value distributions from already-fetched data.")
add_footer(slide, 14)

# 15 — Code behavior
slide = new_slide("How to explain code snippets", "Always connect code to behavior and to what the audience sees.", "Presentation")
add_image_fit(slide, "12_code_to_behavior.png", 0.65, 1.2, 12.05, 4.9)
add_cue(slide, "For every snippet answer three questions: what enters, what happens, and what does the user see?")
add_footer(slide, 15)

# 16 — File map
slide = new_slide("Where the important code lives", "Use this slide if someone asks for implementation details.", "Code map")
add_shape(slide, 0.7, 1.25, 5.95, 4.95, BLUE_LIGHT, BLUE, radius=True)
add_text(slide, "BACKEND", 0.98, 1.55, 5.2, 0.32, size=19, color=BLUE, bold=True)
add_rich_lines(slide, [
    "server/app.py — routes and adapter registry",
    "server/api/config.py — config validation/fingerprint",
    "server/api/contracts.py — canonical models",
    "server/api/configurable.py — fetch/cache/query/fidelity",
    "server/api/common.py — gaps/coverage/matrix math",
    "server/downsample/ — rollups and algorithms",
    "db/cache_schema.sql — generic cache tables",
], 0.95, 2.0, 5.25, 3.65, size=14, bullet=True)
add_shape(slide, 6.85, 1.25, 5.85, 4.95, GREEN_LIGHT, GREEN, radius=True)
add_text(slide, "DATASOURCE + FRONTEND", 7.12, 1.55, 5.3, 0.32, size=19, color=GREEN, bold=True)
add_rich_lines(slide, [
    "datasource/ingest_hlt.py — CSV to source DB",
    "datasource/archive_server.py — /archive/* HTTP API",
    "frontend/app.js — state and API calls",
    "frontend/chart.js — uPlot chart and tooltip",
    "frontend/heatmap.js — Canvas matrix",
    "frontend/histogram.js — distributions",
    "frontend/utils.js — alignment and gap nulls",
], 7.1, 2.0, 5.2, 3.65, size=14, bullet=True)
add_cue(slide, "The generic project never imports the HLT-specific datasource package; HTTP plus config is the boundary.")
add_footer(slide, 16)

# 17 — Guarantees
slide = new_slide("The five guarantees to remember", "These are the strongest claims in the presentation.", "Conclusion")
items = [
    ("1", "Dataset independent", "Any HTTP time-series API can be described by config.json.", BLUE, BLUE_LIGHT),
    ("2", "Source isolated", "Datasource and project cache are separate owners.", RED, RED_LIGHT),
    ("3", "No false gaps", "Missing data remains missing and visible.", GREEN, GREEN_LIGHT),
    ("4", "No silent truncation", "Dense ranges use full SQL bucket aggregation.", ORANGE, ORANGE_LIGHT),
    ("5", "Evidence returned", "Quality, counts, resolution, fidelity, and metrics travel with the series.", PURPLE, PURPLE_LIGHT),
]
for i, (num, title, body, edge, fill) in enumerate(items):
    y = 1.25 + i * 0.98
    add_shape(slide, 0.8, y, 0.62, 0.62, edge, edge, radius=True)
    add_text(slide, num, 0.8, y + 0.11, 0.62, 0.32, size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_shape(slide, 1.65, y, 10.9, 0.62, fill, edge, radius=True)
    add_text(slide, title, 1.92, y + 0.11, 2.7, 0.3, size=17, color=edge, bold=True)
    add_text(slide, body, 4.35, y + 0.11, 7.85, 0.3, size=15, color=INK)
add_cue(slide, "The project is not only fast; it is designed to make the result explainable and honest.")
add_footer(slide, 17)

# 18 — Closing/Q&A
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_shape(slide, 0, 0, W, H, NAVY, NAVY, radius=False)
add_text(slide, "Thank you", 0.9, 1.0, 11.5, 0.75, size=42, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "Anaviz turns arbitrary HTTP time-series data into an evidence-aware visualization.", 1.25, 2.05, 10.8, 0.5, size=21, color=RGBColor(210, 229, 247), align=PP_ALIGN.CENTER)
add_shape(slide, 2.0, 3.2, 9.3, 1.55, RGBColor(34, 77, 119), RGBColor(76, 130, 181), radius=True)
add_text(slide, "Questions?", 2.25, 3.55, 8.8, 0.45, size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "Key closing line: The HLT dataset demonstrates the system, but config.json makes the architecture reusable.", 1.4, 5.45, 10.5, 0.55, size=17, color=RGBColor(216, 232, 246), align=PP_ALIGN.CENTER)
add_text(slide, "ANAVIZ  |  Presentation deck", 0.9, 7.1, 11.5, 0.2, size=9, color=RGBColor(170, 200, 225), align=PP_ALIGN.CENTER)

prs.save(str(OUT))
print(f"Saved {OUT} with {len(prs.slides)} slides")
